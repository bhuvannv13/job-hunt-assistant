"""Excel tracker with a live dashboard.

The workbook has two sheets:

* Jobs: one row per job with role, company, dates, description, match score, trust label, a link,
  the tailored CV file, and columns you fill in (Status, Date applied, Result, Notes) using drop-downs.
* Dashboard: headline numbers and charts built from formulas over the Jobs sheet, so they update
  as soon as you change a status in Excel.

Re-running the tool keeps everything you typed: rows are matched by job ID, new jobs are added and
jobs that have disappeared from the employer's board are marked Closed.
"""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import Job

STATUS = ["Not applied", "Shortlisted", "Applied", "Skipped"]
RESULTS = ["Pending", "Online assessment", "Interview", "Offer", "Rejected", "No response"]
TRUST = ["Verified employer", "Official job board", "Unverified", "Caution"]
COLUMNS = [  # (header, width)
    ("ID", 13), ("Date found", 12), ("Posted", 11), ("Company", 20), ("Role", 38), ("Category", 20), ("Level", 15),
    ("Location", 24), ("Source", 13), ("Trust", 17), ("Match score", 11), ("Matched skills", 34), ("Missing skills", 28),
    ("Description", 60), ("Link", 12), ("Tailored CV", 30), ("Listing", 9), ("Status", 13), ("Date applied", 12),
    ("Result", 17), ("Notes", 30),
]
USER_COLUMNS = ["Status", "Date applied", "Result", "Notes", "Tailored CV"]
TEAL, PALE, INK = "0B6E75", "E6F2F3", "1D2733"
THIN = Side(style="thin", color="D5DDE2")


def _col(name: str) -> str:
    return get_column_letter([c for c, _ in COLUMNS].index(name) + 1)


def _existing_rows(path: Path) -> dict[str, dict]:
    if not path.exists():
        return {}
    ws = load_workbook(path)["Jobs"]
    headers = [c.value for c in ws[1]]
    rows = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, values))
        if record.get("ID"):
            rows[str(record["ID"])] = record
    return rows


def write_tracker(jobs: list[Job], path: str | Path, cv_files: dict[str, str] | None = None, today: date | None = None,
                  overrides: dict[str, dict] | None = None) -> Path:
    """overrides: {job_id: {"Status": ..., "Date applied": ..., "Result": ..., "Notes": ...}} from another tracker
    (the website passes what you recorded in the browser)."""
    path = Path(path)
    today = today or date.today()
    cv_files = cv_files or {}
    old = _existing_rows(path)
    for job_id, values in (overrides or {}).items():
        old[job_id] = {**old.get(job_id, {}), **{k: v for k, v in values.items() if v not in (None, "")}}
    current = {j.id for j in jobs}

    records = []
    for job in jobs:
        prev = old.get(job.id, {})
        records.append({
            "ID": job.id, "Date found": prev.get("Date found") or today, "Posted": job.posted or None,
            "Company": job.company, "Role": job.title, "Category": job.category, "Level": job.level,
            "Location": job.location, "Source": job.source.title(), "Trust": job.trust, "Match score": job.score,
            "Matched skills": ", ".join(job.matched_skills), "Missing skills": ", ".join(job.missing_skills),
            "Description": (job.description[:700] + "…") if len(job.description) > 700 else job.description,
            "Link": job.url, "Tailored CV": cv_files.get(job.id) or prev.get("Tailored CV") or "", "Listing": "Open",
            "Status": prev.get("Status") or "Not applied", "Date applied": prev.get("Date applied"),
            "Result": prev.get("Result") or ("Pending" if prev.get("Status") == "Applied" else None), "Notes": prev.get("Notes"),
        })
    for job_id, prev in old.items():          # keep jobs you tracked that are no longer advertised
        if job_id not in current and prev.get("Role"):
            records.append({**prev, "Listing": "Closed"})
    records.sort(key=lambda r: (r.get("Status") not in ("Applied", "Shortlisted"), -(r.get("Match score") or 0)))

    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    _write_jobs(ws, records)
    _write_dashboard(wb.create_sheet("Dashboard", 0), len(records), today)
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _write_jobs(ws, records):
    headers = [c for c, _ in COLUMNS]
    ws.append(headers)
    for r in records:
        date_found = r.get("Date found")
        if isinstance(date_found, str):
            date_found = date.fromisoformat(date_found[:10])
        posted = r.get("Posted")
        if isinstance(posted, str) and posted:
            posted = date.fromisoformat(posted[:10])
        applied = r.get("Date applied")
        if isinstance(applied, str) and applied:
            try:
                applied = date.fromisoformat(applied[:10])
            except ValueError:
                pass
        ws.append([r.get("ID"), date_found, posted or None, r.get("Company"), r.get("Role"), r.get("Category"), r.get("Level"),
                   r.get("Location"), r.get("Source"), r.get("Trust"), r.get("Match score"), r.get("Matched skills"),
                   r.get("Missing skills"), r.get("Description"), None, r.get("Tailored CV"), r.get("Listing"),
                   r.get("Status"), applied, r.get("Result"), r.get("Notes")])
        link_cell = ws.cell(row=ws.max_row, column=headers.index("Link") + 1)
        if r.get("Link"):
            link_cell.value = "Open advert"
            link_cell.hyperlink = r["Link"]
            link_cell.font = Font(color=TEAL, underline="single")
    n = ws.max_row
    for i, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column_letter in (_col("Description"), _col("Role"),
                                                                                        _col("Matched skills"), _col("Missing skills")))
            cell.border = Border(bottom=THIN)
        ws.row_dimensions[row[0].row].height = 60
    for name in ("Date found", "Posted", "Date applied"):
        for cell in ws[_col(name)][1:]:
            cell.number_format = "dd mmm yyyy"
    ws.freeze_panes = "F2"
    if n >= 2:
        ref = f"A1:{get_column_letter(len(COLUMNS))}{n}"
        table = Table(displayName="Jobs", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(table)
    last = max(n, 2) + 500     # room for rows you add yourself
    for name, options in (("Status", STATUS), ("Result", RESULTS)):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        dv.add(f"{_col(name)}2:{_col(name)}{last}")
        ws.add_data_validation(dv)
    score = _col("Match score")
    ws.conditional_formatting.add(f"{score}2:{score}{last}", ColorScaleRule(start_type="num", start_value=35, start_color="FDECEA",
                                                                            mid_type="num", mid_value=60, mid_color="FFF7D6",
                                                                            end_type="num", end_value=85, end_color="D8F0E3"))
    trust = _col("Trust")
    for label, colour in (("Verified employer", "D8F0E3"), ("Official job board", "E6F2F3"), ("Unverified", "FFF7D6"), ("Caution", "FDECEA")):
        ws.conditional_formatting.add(f"{trust}2:{trust}{last}", CellIsRule(operator="equal", formula=[f'"{label}"'],
                                                                            fill=PatternFill("solid", fgColor=colour)))


def _labels(show_percent=False):
    labels = DataLabelList()
    labels.showVal = not show_percent
    labels.showPercent = show_percent
    labels.showSerName = labels.showCatName = labels.showLegendKey = labels.showLeaderLines = False
    return labels


def _write_dashboard(ws, n_rows, today):
    last = n_rows + 501
    rng = lambda name: f"Jobs!${_col(name)}$2:${_col(name)}${last}"  # noqa: E731
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 17
    ws.column_dimensions["H"].width = 4
    ws.column_dimensions["J"].width = 24
    ws.column_dimensions["K"].width = 12
    ws["B2"] = "Job hunt dashboard"
    ws["B2"].font = Font(size=20, bold=True, color=INK)
    ws["B3"] = f"Built {today:%d %b %Y}. Charts update when you change Status, Date applied or Result on the Jobs sheet."
    ws["B3"].font = Font(size=10, color="566273")

    kpis = [("Jobs tracked", f"=COUNTA({rng('ID')})"),
            ("Applied", f'=COUNTIF({rng("Status")},"Applied")'),
            ("Interviews", f'=COUNTIF({rng("Result")},"Interview")+COUNTIF({rng("Result")},"Offer")'),
            ("Offers", f'=COUNTIF({rng("Result")},"Offer")'),
            ("Response rate", f'=IFERROR((COUNTIF({rng("Result")},"Online assessment")+COUNTIF({rng("Result")},"Interview")'
                              f'+COUNTIF({rng("Result")},"Offer")+COUNTIF({rng("Result")},"Rejected"))/COUNTIF({rng("Status")},"Applied"),0)'),
            ("Avg match score", f"=IFERROR(AVERAGE({rng('Match score')}),0)")]
    for i, (label, formula) in enumerate(kpis):
        col = get_column_letter(2 + i)
        ws[f"{col}5"] = label
        ws[f"{col}5"].font = Font(size=10, color="566273")
        cell = ws[f"{col}6"]
        cell.value = formula
        cell.font = Font(size=22, bold=True, color=TEAL)
        cell.number_format = "0%" if "rate" in label else ("0.0" if "score" in label else "0")
        for r in (5, 6):
            ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=PALE)
            ws[f"{col}{r}"].border = Border(left=Side(style="thick", color="FFFFFF"))
    ws.row_dimensions[6].height = 34

    # Helper tables in columns J-K feed the charts
    def table(title, top, labels, formula_for):
        ws[f"J{top}"], ws[f"K{top}"] = title, "Count"
        ws[f"J{top}"].font = ws[f"K{top}"].font = Font(bold=True, color=INK)
        for i, label in enumerate(labels, start=1):
            ws[f"J{top + i}"] = label
            ws[f"K{top + i}"] = formula_for(label)
        return top, top + len(labels)

    ws["J8"] = "Chart data (calculated from the Jobs sheet)"
    ws["J8"].font = Font(italic=True, size=9, color="566273")
    status_rows = table("Status", 9, STATUS, lambda s: f'=COUNTIF({rng("Status")},"{s}")')
    result_rows = table("Result", 15, RESULTS, lambda s: f'=COUNTIF({rng("Result")},"{s}")')
    from .match import CATEGORIES
    cat_rows = table("Role type", 23, list(CATEGORIES) + ["Other"], lambda s: f'=COUNTIF({rng("Category")},"{s}")')
    trust_rows = table("Trust", 32, TRUST, lambda s: f'=COUNTIF({rng("Trust")},"{s}")')
    weeks = [today - timedelta(days=today.weekday()) - timedelta(weeks=w) for w in range(11, -1, -1)]
    ws["J38"], ws["K38"] = "Week starting", "Applications"
    ws["J38"].font = ws["K38"].font = Font(bold=True, color=INK)
    for i, monday in enumerate(weeks, start=39):
        ws[f"J{i}"] = monday
        ws[f"J{i}"].number_format = "dd mmm"
        ws[f"K{i}"] = f'=COUNTIFS({rng("Date applied")},">="&J{i},{rng("Date applied")},"<"&(J{i}+7))'
    week_rows = (38, 50)

    def refs(rows):
        return (Reference(ws, min_col=11, min_row=rows[0], max_row=rows[1]),
                Reference(ws, min_col=10, min_row=rows[0] + 1, max_row=rows[1]))

    def bar(title, rows, anchor, horizontal=False, colour=TEAL):
        chart = BarChart()
        chart.type = "bar" if horizontal else "col"
        chart.title, chart.style, chart.legend = title, 10, None
        chart.height, chart.width = 7.5, 10.8
        data, cats = refs(rows)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.series[0].graphicalProperties.solidFill = colour
        chart.series[0].graphicalProperties.line.solidFill = colour
        chart.dataLabels = _labels()
        chart.y_axis.majorGridlines = None
        chart.y_axis.delete = chart.x_axis.delete = False
        chart.gapWidth = 60
        ws.add_chart(chart, anchor)

    pie = DoughnutChart()
    pie.title, pie.height, pie.width, pie.holeSize = "Applied or not", 7.5, 10.8, 55
    data, cats = refs(status_rows)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(cats)
    pie.dataLabels = _labels()
    from openpyxl.chart.series import DataPoint
    for idx, colour in enumerate(["C9D3DA", "F2C14E", TEAL, "8A96A3"]):
        point = DataPoint(idx=idx)
        point.graphicalProperties.solidFill = colour
        pie.series[0].dPt.append(point)
    ws.add_chart(pie, "B8")
    bar("Results of applications", result_rows, "E8", colour="2F80ED")
    bar("Jobs by role type", cat_rows, "B24", horizontal=True)
    bar("Jobs by trust check", trust_rows, "E24", horizontal=True, colour="7A5AF8")
    line = LineChart()
    line.title, line.height, line.width, line.legend = "Applications per week", 7.0, 21.8, None
    data, cats = refs(week_rows)
    line.add_data(data, titles_from_data=True)
    line.set_categories(cats)
    line.series[0].graphicalProperties.line.solidFill = TEAL
    line.series[0].graphicalProperties.line.width = 28000
    line.series[0].smooth = False
    line.y_axis.majorGridlines = None
    line.y_axis.delete = line.x_axis.delete = False
    line.x_axis.number_format = "dd mmm"
    ws.add_chart(line, "B40")
