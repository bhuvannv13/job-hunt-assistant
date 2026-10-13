import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from jobhunt.cv import parse
from jobhunt.match import Preferences, rank
from jobhunt.pipeline import run_match
from jobhunt.tailor import cover_letter, load_profile, profile_text, tailor_profile, write_tailored_docx
from jobhunt.tracker import write_tracker
from jobhunt.skills import find_skills
from jobhunt.sources import greenhouse, lever

F = Path(__file__).parent / "fixtures"
PROFILE = Path(__file__).parent.parent / "examples" / "profile.example.yaml"


def ranked_jobs():
    load = lambda n: json.loads((F / f"{n}.json").read_text())  # noqa: E731
    jobs = greenhouse.parse(load("greenhouse"), "Stripe") + lever.parse(load("lever"), "Palantir")
    cv = parse((F / "cv.txt").read_text())
    return rank(jobs, cv.text, cv.skills, Preferences(min_score=0))[0]


def test_tailoring_only_uses_profile_content(tmp_path):
    profile = load_profile(PROFILE)
    ml_job = next(j for j in ranked_jobs() if "Machine Learning" in j.title)
    tailored = tailor_profile(profile, ml_job)
    assert tailored["projects"][0]["name"] == "Image classifier"           # most relevant first
    assert tailored["summary"] == profile["summaries"]["AI / ML engineer"]
    known = set(find_skills(profile_text(profile)))
    assert set(tailored["key_skills"]) <= known                              # nothing invented
    all_bullets = {b for item in profile["experience"] + profile["projects"] for b in item["bullets"]}
    for item in tailored["experience"] + tailored["projects"]:
        assert set(item["bullets"]) <= all_bullets
    path = write_tailored_docx(tailored, tmp_path / "cv.docx")
    assert path.exists() and path.stat().st_size > 5000
    letter = cover_letter(profile, ml_job, tailored)
    assert ml_job.company in letter and ml_job.title in letter and "Alex Example" in letter


def test_tracker_keeps_user_edits(tmp_path):
    jobs = ranked_jobs()
    path = tmp_path / "tracker.xlsx"
    write_tracker(jobs, path, today=date(2026, 10, 1))
    wb = load_workbook(path)
    assert wb.sheetnames == ["Dashboard", "Jobs"]
    ws = wb["Jobs"]
    headers = [c.value for c in ws[1]]
    status_col = headers.index("Status") + 1
    ws.cell(row=2, column=status_col, value="Applied")
    ws.cell(row=2, column=headers.index("Result") + 1, value="Interview")
    edited_id = ws.cell(row=2, column=1).value
    wb.save(path)
    write_tracker(jobs[1:] if jobs[0].id != edited_id else jobs[1:], path, today=date(2026, 10, 8))
    ws = load_workbook(path)["Jobs"]
    rows = {r[0]: dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows[edited_id]["Status"] == "Applied" and rows[edited_id]["Result"] == "Interview"
    assert len(load_workbook(path)["Dashboard"]._charts) == 5


def test_full_match_run(tmp_path):
    load = lambda n: json.loads((F / f"{n}.json").read_text())  # noqa: E731
    jobs = greenhouse.parse(load("greenhouse"), "Stripe") + lever.parse(load("lever"), "Palantir")
    profile = load_profile(PROFILE)
    text = profile_text(profile)
    summary = run_match(jobs, text, find_skills(text), Preferences(min_score=0), tmp_path, profile=profile, tailor_top=2)
    assert summary["tailored"] == 2
    assert len(list((tmp_path / "cvs").glob("*.docx"))) == 2
    assert len(list((tmp_path / "cover_letters").glob("*.docx"))) == 2
    assert (tmp_path / "job_tracker.xlsx").exists() and (tmp_path / "shortlist.md").exists()
