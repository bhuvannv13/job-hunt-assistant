"""Read a CV (PDF, DOCX, TXT or Markdown) and split it into sections."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from .skills import find_skills

SECTION_NAMES = {
    "profile": ["profile", "summary", "about me", "personal statement", "objective", "professional summary"],
    "education": ["education", "academic background", "qualifications"],
    "skills": ["skills", "technical skills", "core skills", "key skills", "competencies"],
    "experience": ["experience", "work experience", "employment", "professional experience", "work history"],
    "projects": ["projects", "personal projects", "selected projects", "academic projects", "other projects"],
    "awards": ["awards", "achievements", "awards and activities", "certifications", "honours", "activities"],
}
_HEADING = {alias: key for key, aliases in SECTION_NAMES.items() for alias in aliases}
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"\+?\d[\d ()-]{7,}\d")


@dataclass
class CV:
    text: str
    sections: dict[str, str] = field(default_factory=dict)
    skills: list[str] = field(default_factory=list)
    name: str = ""
    email: str = ""
    phone: str = ""

    def bullets(self, section: str) -> list[str]:
        """Lines in a section that look like individual points."""
        lines = [l.strip() for l in self.sections.get(section, "").split("\n") if l.strip()]
        marked = [l for l in lines if l[0] in "•▪◦-*–·"]
        chosen = marked or lines
        return [l.lstrip("•▪◦-*–· ").strip() for l in chosen if len(l) > 25]


def read_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            import pdfplumber
            with pdfplumber.open(str(path)) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        except ImportError:
            from pypdf import PdfReader
            return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if suffix in (".txt", ".md", ".text"):
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported CV format: {suffix} (use PDF, DOCX, TXT or MD)")


def split_sections(text: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw in text.split("\n"):
        line = raw.strip()
        key = _HEADING.get(re.sub(r"[^a-z ]", "", line.lower()).strip())
        if key and len(line) < 40:
            current = key
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(raw)
    return {k: "\n".join(v).strip() for k, v in sections.items() if "\n".join(v).strip()}


def parse(text: str) -> CV:
    sections = split_sections(text)
    header_lines = [l.strip() for l in sections.get("header", "").split("\n") if l.strip()]
    email = EMAIL.search(text)
    phone = PHONE.search(sections.get("header", text[:600]))
    return CV(text=text, sections=sections, skills=find_skills(text),
              name=header_lines[0] if header_lines else "",
              email=email.group(0) if email else "", phone=phone.group(0).strip() if phone else "")


def load(path: str | Path) -> CV:
    return parse(read_text(path))
